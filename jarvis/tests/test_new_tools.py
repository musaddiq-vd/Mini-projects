import platform
import sys

import pytest

windows_only = pytest.mark.skipif(platform.system() != "Windows", reason="Windows only")


def test_all_tools_load(registry):
    assert not registry.errors, registry.errors
    for name in ["close_app", "list_windows", "window_control", "press_keys", "type_text", "brightness",
                 "dark_mode", "wifi", "look_at_screen", "read_clipboard", "read_document", "send_whatsapp",
                 "draft_email", "game_launcher", "find_flights", "write_code", "watch_topic", "autostart"]:
        assert name in registry.tools, name


@windows_only
def test_close_app_closes_matching_window_or_tab(registry, ctx, monkeypatch):
    from jarvis.tools import _win, windows

    wins = [
        _win.Window(1, "Untitled - Notepad", "Notepad.exe", 10),
        _win.Window(2, "YouTube - Google Chrome", "chrome.exe", 20),
    ]
    closed, focused, keys = [], [], []
    monkeypatch.setattr(_win, "list_windows", lambda: wins)
    monkeypatch.setattr(_win, "close", closed.append)
    monkeypatch.setattr(_win, "focus", lambda h: focused.append(h) or True)
    monkeypatch.setattr(windows, "_press", lambda *k: keys.append(k))

    assert "Closed Notepad" in registry.call("close_app", {"name": "notepad"}, ctx)["result"]
    assert closed == [1]

    out = registry.call("close_app", {"name": "YouTube"}, ctx)["result"]
    assert "tab" in out and focused == [2] and keys == [("ctrl", "w")]
    assert closed == [1]  # the browser window itself was not closed

    registry.call("close_app", {"name": "chrome"}, ctx)
    assert closed == [1, 2]
    assert "error" in registry.call("close_app", {"name": "photoshop"}, ctx)


def test_press_keys_validation(registry, ctx, monkeypatch):
    from jarvis.tools import windows

    pressed = []
    monkeypatch.setattr(windows, "_press", lambda *k: pressed.append(k))
    assert registry.call("press_keys", {"keys": "Ctrl + Shift + T"}, ctx)["result"] == "Pressed ctrl+shift+t."
    assert "error" in registry.call("press_keys", {"keys": "ctrl+;rm -rf"}, ctx)
    assert pressed == [("ctrl", "shift", "t")]


def test_brightness_with_undo(registry, ctx, monkeypatch):
    from jarvis.tools import settings_tools as st

    level = {"v": 50}
    monkeypatch.setattr(st, "get_brightness", lambda: level["v"])
    monkeypatch.setattr(st, "set_brightness_level", lambda v: level.update(v=v) or True)
    registry.call("brightness", {"level": "up"}, ctx)
    assert level["v"] == 60
    registry.call("brightness", {"level": "25%"}, ctx)
    assert level["v"] == 25
    assert "error" in registry.call("brightness", {"level": "max; shutdown"}, ctx)
    registry.call("undo", {}, ctx)
    assert level["v"] == 60


def test_whatsapp_needs_number_and_confirmation(registry, ctx, answers, monkeypatch):
    from jarvis.tools import messaging

    opened = []
    monkeypatch.setattr(messaging, "_press_send_when_ready", lambda: "Message sent.")
    monkeypatch.setattr(messaging.os, "startfile", opened.append, raising=False)
    monkeypatch.setattr(messaging.webbrowser, "open", opened.append)

    assert "don't have a number" in registry.call("send_whatsapp", {"contact": "Ali", "message": "hi"}, ctx)["error"]
    ctx.memory.remember("contact_ali", "+91 98765 43210")
    answers.append(False)
    assert "Cancelled" in registry.call("send_whatsapp", {"contact": "Ali", "message": "hi"}, ctx)["result"]
    assert not opened
    answers.append(True)
    assert registry.call("send_whatsapp", {"contact": "Ali", "message": "kal milte hain"}, ctx)["result"] == "Message sent."
    assert "phone=919876543210" in opened[0] and "kal%20milte%20hain" in opened[0]


def test_watch_topic_and_background_check(ctx, registry):
    from types import SimpleNamespace

    from jarvis.background import BackgroundServices
    from jarvis.events import Events
    from jarvis.state import StateStore

    ctx.state = StateStore(ctx.memory.path.parent / "state.json")
    registry.call("watch_topic", {"action": "add", "topic": "Bitcoin"}, ctx)
    registry.call("watch_topic", {"action": "add", "topic": "bitcoin"}, ctx)
    assert ctx.state.get("watch_topics") == ["Bitcoin"]

    said = []
    app = SimpleNamespace(settings=ctx.settings, state=ctx.state, memory=ctx.memory, events=Events(), ctx=ctx)
    bg = BackgroundServices(app, said.append)
    headlines = [[{"title": "A"}, {"title": "B"}], [{"title": "C"}, {"title": "A"}]]
    assert bg.check_topics(search=lambda q: headlines[0]) == []  # first check only records
    assert bg.check_topics(search=lambda q: headlines[1]) == ["Bitcoin: C"]
    assert "Bitcoin: C" in said[0]


def test_briefing_once_per_day(ctx):
    from types import SimpleNamespace

    from jarvis.background import BackgroundServices
    from jarvis.events import Events
    from jarvis.state import StateStore

    ctx.state = StateStore(ctx.memory.path.parent / "state.json")
    ctx.settings.city = "Pune"
    said = []
    app = SimpleNamespace(settings=ctx.settings, state=ctx.state, memory=ctx.memory, events=Events(), ctx=ctx)
    bg = BackgroundServices(app, said.append)
    bg.on_connected()
    bg.on_connected()
    assert len(said) == 1 and "Pune" in said[0] and "news" in said[0]


def test_documents_docx_and_xlsx(registry, ctx, sandbox):
    import docx
    import openpyxl

    d = docx.Document()
    d.add_paragraph("Quarterly report: sales up 20%")
    d.save(sandbox / "report.docx")
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Marks"])
    wb.active.append(["Musaddiq", 95])
    wb.save(sandbox / "marks.xlsx")
    assert "sales up 20%" in registry.call("read_document", {"path": str(sandbox / "report.docx")}, ctx)["content"]
    assert "Musaddiq | 95" in registry.call("read_document", {"path": str(sandbox / "marks.xlsx")}, ctx)["content"]


def test_write_code_stays_inside_folder(registry, ctx, sandbox, monkeypatch):
    import json

    from jarvis.tools import code_helper

    reply = json.dumps({"files": [{"path": "main.py", "content": "print('hi')"},
                                  {"path": "../escape.py", "content": "bad"}], "run": "python main.py"})
    ctx._llm = type("L", (), {"ask": lambda self, p, system=None: "```json\n" + reply + "\n```"})()
    monkeypatch.setattr(code_helper, "_open_in_editor", lambda p: True)
    out = registry.call("write_code", {"description": "hello", "folder": str(sandbox / "proj")}, ctx)
    assert out["files"] == ["main.py"]
    assert (sandbox / "proj" / "main.py").read_text() == "print('hi')"
    assert not (sandbox / "escape.py").exists()


def test_events_bus_isolates_bad_handlers(capsys):
    from jarvis.events import Events

    bus, got = Events(), []
    bus.subscribe(lambda k, d: 1 / 0)
    bus.subscribe(lambda k, d: got.append((k, d)))
    bus.emit("log", text="x")
    assert got == [("log", {"text": "x"})]


def test_window_builds_offscreen(monkeypatch, tmp_path):
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication

    from jarvis.app import create_app
    from jarvis.config import Settings
    from jarvis.ui.window import Bridge, MainWindow

    qt = QApplication.instance() or QApplication(sys.argv)
    bridge = Bridge()
    app = create_app(Settings(theme_color="#ff3df2"))
    w = MainWindow(app, bridge)
    app.events.emit("state", state="listening")
    app.events.emit("user", text="close youtube")
    app.events.emit("tool", name="close_app", args={"name": "youtube"}, result={"result": "Closed tab."})
    app.events.emit("assistant", text="Done!")
    qt.processEvents()
    assert "close youtube" in w.chat.toPlainText() and "Closed tab." in w.chat.toPlainText()
    assert w.state_label.text() == "LISTENING"
    w.grab().save(str(tmp_path / "window.png"))


def test_computer_sleep_needs_confirmation(registry, ctx, answers, monkeypatch):
    import jarvis.tools.system as system

    slept, launched = [], []
    monkeypatch.setattr(system, "_windows_sleep", lambda: slept.append(1))
    monkeypatch.setattr(system.subprocess, "Popen", launched.append)
    monkeypatch.setattr(system.threading, "Timer", lambda d, fn: type("T", (), {"start": lambda self: fn()})())
    for action in ("sleep", "shutdown", "restart"):
        answers.append(False)
        assert "Cancelled" in registry.call("power", {"action": action}, ctx)["result"]
    assert not slept and not launched
    answers.append(True)
    registry.call("power", {"action": "sleep"}, ctx)
    assert slept or launched


def test_assistant_sleep_never_touches_power(registry, ctx, monkeypatch):
    import jarvis.tools.assistant_tools as at
    import jarvis.tools.system as system

    monkeypatch.setattr(at.threading, "Timer", lambda d, fn: type("T", (), {"start": lambda self: fn()})())
    monkeypatch.setattr(system.subprocess, "Popen", lambda *a: (_ for _ in ()).throw(AssertionError("power used")))
    calls = []
    ctx.assistant = type("A", (), {"sleep": lambda self: calls.append("sleep"), "wake": None})()
    out = registry.call("assistant_sleep", {}, ctx)
    assert calls == ["sleep"] and "Going to sleep" in out["result"]
    ctx.assistant = None
    assert "error" in registry.call("assistant_sleep", {}, ctx)


def test_sleep_rule_in_prompt(ctx):
    from jarvis.app import system_prompt

    text = system_prompt(ctx.settings, ctx.memory)
    assert "assistant_sleep" in text and "SLEEP RULE" in text


def test_live_restart_keeps_awake_and_app_messages_dont_wake(ctx):
    from jarvis.app import create_app
    from jarvis.config import Settings
    from jarvis.live import LiveAssistant

    app = create_app(Settings(api_key="test-key"))
    a = LiveAssistant(app)
    assert app.ctx.assistant is a
    a.restart()
    assert a._reload_wake is False  # no wake word -> nothing to reload, stays awake
    a.sleep()
    a.send_system("(system: hardware alert) battery low")
    assert a.awake is False  # alerts speak but don't open the microphone
    a.send_text("hello")
    assert a.awake is True

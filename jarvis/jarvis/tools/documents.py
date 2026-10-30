"""Read PDF, Word, Excel, PowerPoint, CSV and text files so the model can summarise or answer."""

from __future__ import annotations

import os
from pathlib import Path

from jarvis.tools import tool
from jarvis.tools.files import PathNotAllowed, resolve

MAX_CHARS = 30_000
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".py", ".js", ".html", ".css", ".log", ".xml", ".yaml", ".yml"}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join(f"[page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages[:60]))
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if suffix in {".xlsx", ".xlsm"}:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:10]:
            out.append(f"[sheet {ws.title}]")
            for row in ws.iter_rows(max_row=300, values_only=True):
                if any(v is not None for v in row):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    if suffix == ".pptx":
        from pptx import Presentation

        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            out.append(f"[slide {i}] " + " / ".join(texts))
        return "\n".join(out)
    if suffix in TEXT_SUFFIXES or not suffix:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Can't read {suffix} files.")


@tool(
    "read_document",
    "Read a PDF, Word, Excel, PowerPoint, CSV or text file to summarise it or answer questions about it. "
    "Paths can start with desktop/, documents/, downloads/.",
    {"path": {"type": "string"}},
    required=["path"],
    slow=True,
)
def read_document(ctx, path: str):
    try:
        file = resolve(ctx, path)
    except PathNotAllowed as exc:
        return {"error": str(exc)}
    if not file.exists():
        matches = _find_by_name(ctx, Path(path).name)
        if len(matches) == 1:
            file = matches[0]
        elif matches:
            return {"error": "Several files match; which one?", "matches": [str(m) for m in matches[:10]]}
        else:
            return {"error": f"File not found: {file}"}
    try:
        text = extract_text(file)
    except ValueError as exc:
        return {"error": str(exc)}
    return {"path": str(file), "content": text[:MAX_CHARS], "truncated": len(text) > MAX_CHARS}


def _find_by_name(ctx, name: str) -> list[Path]:
    """Look for a file by name in the usual folders (shallow, fast)."""
    from jarvis.tools.files import SHORTCUTS

    name = name.lower()
    found = []
    for key in ("desktop", "documents", "downloads"):
        folder = SHORTCUTS[key]
        if not folder.is_dir():
            continue
        base_depth = len(folder.parts)
        for root, dirs, files in os.walk(folder):
            if len(Path(root).parts) - base_depth >= 3:
                dirs.clear()  # don't go deeper than 3 levels
            dirs[:] = [d for d in dirs if not d.startswith(".") and d not in {"node_modules", "venv", ".venv"}]
            for f in files:
                if name in f.lower():
                    found.append(Path(root) / f)
                    if len(found) > 20:
                        return found
    return found
